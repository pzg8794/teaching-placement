from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


@dataclass(frozen=True)
class Theme:
    ink: RGBColor = RGBColor.from_string("172033")
    paper: RGBColor = RGBColor.from_string("F7FBFF")
    sky: RGBColor = RGBColor.from_string("D9ECFF")
    leaf: RGBColor = RGBColor.from_string("BFE8A5")
    sun: RGBColor = RGBColor.from_string("FFD66B")
    coral: RGBColor = RGBColor.from_string("FFB7A5")
    blue: RGBColor = RGBColor.from_string("2563EB")
    green: RGBColor = RGBColor.from_string("2F7D32")
    gold: RGBColor = RGBColor.from_string("9A6700")
    white: RGBColor = RGBColor.from_string("FFFFFF")


THEME = Theme()
OUT = Path(__file__).resolve().parent / "k-2-ai-build-ideas-with-magicschool.pptx"


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def add_shape(slide, shape_type, x, y, w, h, fill=None, line=None, radius=True):
    shp = slide.shapes.add_shape(
        shape_type,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(1.25)
    return shp


def add_round(slide, x, y, w, h, fill, line=None):
    return add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h, fill, line)


def add_text(slide, x, y, w, h, text, size=32, bold=False, color=None, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.04)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color or THEME.ink
    return box


def add_title(slide, title):
    add_text(slide, 0.55, 0.28, 12.2, 0.72, title, size=34, bold=True)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0.55, 1.05, 12.2, 0.06, THEME.sun)


def add_simple_bullets(slide, x, y, w, h, items, size=31):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.04)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(7)
        p.alignment = PP_ALIGN.LEFT
        r = p.add_run()
        r.text = f"- {item}"
        r.font.name = "Arial"
        r.font.size = Pt(size)
        r.font.bold = True
        r.font.color.rgb = THEME.ink
    return box


def add_prompt_card(slide, prompt):
    add_round(slide, 0.85, 1.65, 11.65, 3.7, THEME.white, THEME.blue)
    add_text(slide, 1.15, 1.95, 11.05, 2.95, prompt, size=34, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, 1.2, 5.65, 10.95, 0.5, "Copy the words. Fill in the blank.", size=24, bold=True, color=THEME.blue, align=PP_ALIGN.CENTER)


def add_material_icon(slide, x, y, label, fill, icon_type):
    add_round(slide, x, y, 2.2, 1.55, fill, THEME.ink)
    if icon_type == "circles":
        for dx, dy in [(0.45, 0.42), (0.95, 0.42), (1.45, 0.42), (0.7, 0.85), (1.2, 0.85)]:
            add_shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, x + dx, y + dy, 0.32, 0.32, THEME.white, THEME.ink)
    elif icon_type == "blocks":
        for dx, dy in [(0.4, 0.38), (0.92, 0.38), (1.44, 0.38), (0.65, 0.82), (1.17, 0.82)]:
            add_shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, x + dx, y + dy, 0.42, 0.34, THEME.white, THEME.ink)
    elif icon_type == "sticks":
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, x + 0.45, y + 0.45, 1.28, 0.12, THEME.white, THEME.ink)
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, x + 0.45, y + 0.83, 1.28, 0.12, THEME.white, THEME.ink)
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, x + 0.82, y + 0.25, 0.12, 1.0, THEME.white, THEME.ink)
    else:
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, x + 0.48, y + 0.44, 1.25, 0.22, THEME.white, THEME.ink)
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, x + 0.48, y + 0.82, 1.25, 0.22, THEME.white, THEME.ink)
    add_text(slide, x, y + 1.08, 2.2, 0.42, label, size=17, bold=True, align=PP_ALIGN.CENTER)


def add_step_slide(prs, title, big, small=None, fill=None):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_shape(s, MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, 13.333, 7.5, fill or THEME.paper)
    add_title(s, title)
    add_round(s, 0.9, 1.75, 11.55, 3.85, THEME.white, THEME.ink)
    add_text(s, 1.3, 2.15, 10.75, 2.85, big, size=44, bold=True, align=PP_ALIGN.CENTER)
    if small:
        add_text(s, 1.0, 6.05, 11.3, 0.55, small, size=25, bold=True, color=THEME.blue, align=PP_ALIGN.CENTER)
    return s


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # 1
    s = prs.slides.add_slide(blank)
    add_shape(s, MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, 13.333, 7.5, THEME.paper)
    add_round(s, 0.55, 0.55, 12.22, 6.4, THEME.sky, None)
    add_text(s, 0.95, 1.15, 7.7, 1.0, "AI Build Ideas", size=50, bold=True)
    add_text(s, 1.0, 2.1, 7.3, 0.75, "Ask. Choose. Build.", size=36, bold=True, color=THEME.blue)
    add_round(s, 0.9, 3.05, 7.7, 1.95, THEME.white, THEME.blue)
    add_text(s, 1.2, 3.35, 7.1, 1.25, "AI helps us get ideas.\nWe choose what to build.", size=33, bold=True, align=PP_ALIGN.CENTER)
    add_material_icon(s, 9.15, 1.15, "materials", THEME.leaf, "blocks")
    add_material_icon(s, 9.95, 3.0, "build", THEME.sun, "sticks")

    # 2
    s = prs.slides.add_slide(blank)
    add_shape(s, MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, 13.333, 7.5, THEME.paper)
    add_title(s, "Today We Will")
    add_simple_bullets(s, 1.25, 1.55, 10.9, 4.9, [
        "ask AI for ideas",
        "choose one idea",
        "build with our material",
        "share what we made",
    ], size=34)

    # 3
    s = prs.slides.add_slide(blank)
    add_shape(s, MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, 13.333, 7.5, THEME.paper)
    add_title(s, "What Is on Your Table?")
    add_material_icon(s, 0.95, 1.55, "Zoobs", THEME.leaf, "circles")
    add_material_icon(s, 3.45, 1.55, "LEGO", THEME.sun, "blocks")
    add_material_icon(s, 5.95, 1.55, "Strawbees", THEME.coral, "sticks")
    add_material_icon(s, 8.45, 1.55, "KEVA", THEME.sky, "planks")
    add_material_icon(s, 10.95, 1.55, "Other", rgb("E5E7EB"), "blocks")
    add_text(s, 1.0, 4.1, 11.3, 1.1, "Say: I have ______.", size=44, bold=True, align=PP_ALIGN.CENTER)

    # 4
    s = prs.slides.add_slide(blank)
    add_shape(s, MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, 13.333, 7.5, THEME.paper)
    add_title(s, "Our AI Rules")
    add_simple_bullets(s, 0.95, 1.45, 11.4, 4.75, [
        "use the class link in Seesaw",
        "do not type your name",
        "do not type private information",
        "AI gives ideas",
        "we choose what to build",
    ], size=29)

    # 5-6
    add_step_slide(prs, "Step 1", "Open Seesaw.\nTap the MagicSchool link.", "Then wait for the teacher.")
    add_step_slide(prs, "Step 2", "Look at your table.\nSay: I have ____.", "Use the word for your building material.")

    # 7
    s = prs.slides.add_slide(blank)
    add_shape(s, MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, 13.333, 7.5, THEME.paper)
    add_title(s, "Copy This Prompt")
    add_prompt_card(s, "I have ____ at my table.\nPlease give me 3 easy things I can build.\nUse simple words.")

    # 8
    s = prs.slides.add_slide(blank)
    add_shape(s, MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, 13.333, 7.5, THEME.paper)
    add_title(s, "If the AI Answer Is Too Hard")
    add_prompt_card(s, "Make it easier.\nUse short words.\nGive me 3 simple ideas.")

    # 9
    s = prs.slides.add_slide(blank)
    add_shape(s, MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, 13.333, 7.5, THEME.paper)
    add_title(s, "Pick One Idea")
    add_simple_bullets(s, 1.2, 1.8, 10.4, 3.4, [
        "Can I build it?",
        "Do we like it?",
        "Can we start now?",
    ], size=38)
    add_text(s, 1.2, 5.65, 10.8, 0.7, "Choose one. It can be simple.", size=28, bold=True, color=THEME.green, align=PP_ALIGN.CENTER)

    # 10
    s = prs.slides.add_slide(blank)
    add_shape(s, MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, 13.333, 7.5, THEME.paper)
    add_title(s, "Build Time")
    steps = [("1", "Choose", THEME.sun), ("2", "Build", THEME.leaf), ("3", "Fix", THEME.coral), ("4", "Share", THEME.sky)]
    x = 0.9
    for num, word, fill in steps:
        add_round(s, x, 2.0, 2.85, 2.05, fill, THEME.ink)
        add_text(s, x, 2.25, 2.85, 0.55, num, size=34, bold=True, align=PP_ALIGN.CENTER)
        add_text(s, x, 2.9, 2.85, 0.55, word, size=28, bold=True, align=PP_ALIGN.CENTER)
        x += 3.05

    # 11
    s = prs.slides.add_slide(blank)
    add_shape(s, MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, 13.333, 7.5, THEME.paper)
    add_title(s, "If You Get Stuck")
    add_prompt_card(s, "Give me a new idea using ____.\nMake it easier.")

    # 12
    s = prs.slides.add_slide(blank)
    add_shape(s, MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, 13.333, 7.5, THEME.paper)
    add_title(s, "Share")
    add_prompt_card(s, "We built ____.\nAI helped us ____.\nWe changed ____.")

    # 13
    s = prs.slides.add_slide(blank)
    add_shape(s, MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, 13.333, 7.5, THEME.paper)
    add_title(s, "Teacher Notes: NYS CSDF")
    add_text(s, 0.85, 1.45, 12.0, 0.55, "Aligned K-1 standards with Grade 2 bridge:", size=24, bold=True)
    add_simple_bullets(s, 0.9, 2.05, 11.9, 4.3, [
        "K-1.IC.1, K-1.IC.2, K-1.IC.6",
        "K-1.CT.4, K-1.CT.6, K-1.CT.10",
        "K-1.CY.1",
        "K-1.DL.2, K-1.DL.3, K-1.DL.4, K-1.DL.7",
        "Grade 2 bridge: 2-3.IC.1, 2-3.CY.1, 2-3.DL.3, 2-3.NSD.3",
    ], size=20)

    # 14
    s = prs.slides.add_slide(blank)
    add_shape(s, MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, 13.333, 7.5, THEME.paper)
    add_title(s, "Teacher Notes: Inclusion")
    add_simple_bullets(s, 0.9, 1.45, 11.9, 5.25, [
        "Use visual directions and repeated sentence frames.",
        "Students may copy, say, or partner-type prompts.",
        "Roles: typer, reader, builder, sharer.",
        "Accept a build, photo, pointing, oral share, or partner share.",
        "Short predictable prompts support early readers and multilingual learners.",
    ], size=21)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
