from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


OUT = Path("lesson2-3-wrap-up-slides-editable.pptx")

ACCENT1 = RGBColor.from_string("0F766E")
ACCENT2 = RGBColor.from_string("99F6E4")
ACCENT3 = RGBColor.from_string("F59E0B")
ACCENT4 = RGBColor.from_string("134E4A")
PAPER = RGBColor.from_string("F7FEFC")
INK = RGBColor.from_string("1F2937")
WHITE = RGBColor(255, 255, 255)


PROMPTS = {
    "recover": "My topic is ____. Help me make a simple project guide with 3 parts. Each part should be short and connected to computer science.",
    "build": "My topic is ____. Part __ of my guide is ____. Give me 3 short bullet notes I can use on one slide. Keep it clear for Grade 4.",
    "check": "My topic is ____. My guide is ____. Here is my project text: ____. Tell me what matches, what is missing, and what I should fix first.",
    "finish": "Help me write a short final reflection. I used AI to ____. I made my own choices by ____. Keep it 2 sentences.",
}


def add_rect(slide, x, y, w, h, fill=WHITE, line=None, rounded=True):
    shape_type = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if rounded else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    return shape


def add_text(slide, x, y, w, h, text, size=24, bold=False, color=INK, align=PP_ALIGN.LEFT, font="Arial"):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_title(slide, title):
    add_text(slide, 0.65, 0.35, 12.0, 0.6, title, size=34, bold=True)


def add_bullets(slide, x, y, w, h, items, size=28):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        r1 = p.add_run()
        r1.text = "> "
        r1.font.name = "Arial"
        r1.font.size = Pt(size)
        r1.font.bold = True
        r1.font.color.rgb = ACCENT1
        r2 = p.add_run()
        r2.text = item
        r2.font.name = "Arial"
        r2.font.size = Pt(size)
        r2.font.color.rgb = INK
    return box


def add_block(slide, x, y, w, h, header, body, body_size=26):
    add_rect(slide, x, y, w, h, fill=WHITE, line=ACCENT4, rounded=True)
    add_text(slide, x + 0.25, y + 0.16, w - 0.5, 0.35, header, size=18, bold=True, color=ACCENT4)
    if isinstance(body, list):
        add_bullets(slide, x + 0.35, y + 0.65, w - 0.7, h - 0.75, body, size=body_size)
    else:
        add_text(slide, x + 0.35, y + 0.72, w - 0.7, h - 0.85, body, size=body_size)


def add_prompt(slide, prompt_id, y, label):
    add_rect(slide, 0.8, y - 0.04, 1.05, 0.42, fill=RGBColor.from_string("CCFBF1"), line=ACCENT1, rounded=True)
    add_text(slide, 0.85, y, 1.0, 0.35, "Copy", size=14, bold=True, color=ACCENT4, align=PP_ALIGN.CENTER)
    add_text(slide, 2.0, y - 0.03, 10.55, 0.85, PROMPTS[prompt_id], size=20, color=INK, font="Courier New")
    add_text(slide, 2.0, y - 0.43, 10.55, 0.3, label, size=13, bold=True, color=ACCENT4)


def add_slide(prs, title):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(slide, 0, 0, 13.333, 7.5, fill=PAPER, line=None, rounded=False)
    add_title(slide, title)
    return slide


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(slide, 0, 0, 13.333, 7.5, fill=PAPER, line=None, rounded=False)
    add_rect(slide, 0.45, 0.5, 12.45, 6.45, fill=RGBColor.from_string("E6FFFB"), line=None, rounded=True)
    add_rect(slide, 0.65, 6.72, 12.05, 0.08, fill=ACCENT3, line=None, rounded=False)
    add_rect(slide, 0.65, 0.72, 12.05, 0.08, fill=ACCENT2, line=None, rounded=False)
    add_text(slide, 0.85, 1.15, 11.7, 0.85, "Lesson 2-3: Build, Check, Finish", size=42, bold=True)
    add_text(slide, 0.85, 2.0, 11.7, 0.45, "Wrap up the AI project in 4 simple steps.", size=26, color=ACCENT4)
    for i, label in enumerate(["1 Recover", "2 Build", "3 Check", "4 Finish"]):
        x = 1.15 + i * 2.95
        fill = [RGBColor.from_string("FDE68A"), RGBColor.from_string("CCFBF1"), RGBColor.from_string("D1FAE5"), RGBColor.from_string("FDE68A")][i]
        border = [ACCENT3, ACCENT1, ACCENT1, ACCENT3][i]
        add_rect(slide, x, 3.25, 2.35, 1.05, fill=fill, line=border, rounded=True)
        add_text(slide, x, 3.5, 2.35, 0.45, label, size=22, bold=True, align=PP_ALIGN.CENTER)
    add_block(slide, 1.0, 5.0, 11.3, 1.0, "Big idea", "AI can help us finish, but we make the choices.", body_size=26)

    slide = add_slide(prs, "Today You Need")
    add_block(slide, 0.9, 1.25, 11.5, 4.15, "By the end, you should have:", ["a topic", "3 project parts or slides", "one AI check", "one final reflection"], body_size=28)
    add_text(slide, 0.9, 6.0, 11.5, 0.5, "Build first. Check second. Finish third.", size=28, bold=True, color=ACCENT4, align=PP_ALIGN.CENTER)

    slide = add_slide(prs, "Step 1: Recover")
    add_block(slide, 0.9, 1.25, 11.5, 1.2, "If you already have a topic and guide", "Go straight to Step 2.", body_size=28)
    add_block(slide, 0.9, 2.9, 11.5, 3.2, "If you are missing your guide", "", body_size=22)
    add_prompt(slide, "recover", 3.95, "Recover prompt")

    slide = add_slide(prs, "Step 2: Build One Slide")
    add_block(slide, 0.9, 1.25, 11.5, 3.5, "Use this for each part", "", body_size=22)
    add_prompt(slide, "build", 2.25, "Build prompt")
    add_block(slide, 0.9, 5.15, 11.5, 1.0, "Build rule", "Use the notes to make your own slide. Read first, then decide.", body_size=24)

    slide = add_slide(prs, "Simple Example")
    add_block(slide, 0.75, 1.3, 5.95, 4.6, "Topic and guide", "Topic: robots help animals\n\nPart 1: what robots are\n\nPart 2: how robots help\n\nPart 3: why it matters", body_size=21)
    add_block(slide, 6.95, 1.3, 5.65, 4.6, "One slide", "Title: How robots help\n\nRobots can use sensors.\n\nRobots can follow code.\n\nRobots can do jobs safely.", body_size=21)

    slide = add_slide(prs, "Step 3: Check Your Project")
    add_block(slide, 0.9, 1.25, 11.5, 3.65, "Use this after you build", "", body_size=22)
    add_prompt(slide, "check", 2.2, "Check prompt")
    add_block(slide, 0.9, 5.25, 11.5, 1.0, "Check rule", "Fix one thing first. Keep it simple.", body_size=25)

    slide = add_slide(prs, "Step 4: Finish")
    add_block(slide, 0.9, 1.25, 11.5, 3.35, "Add your final reflection", "", body_size=22)
    add_prompt(slide, "finish", 2.25, "Finish prompt")
    add_block(slide, 0.9, 5.0, 11.5, 1.0, "Finish rule", "Your reflection should sound like you.", body_size=26)

    slide = add_slide(prs, "Before You Turn It In")
    add_block(slide, 0.9, 1.25, 11.5, 4.65, "Final checklist", ["My topic is clear.", "I have 3 parts or slides.", "I checked my project with AI.", "I fixed one thing.", "I wrote my reflection."], body_size=27)

    slide = add_slide(prs, "Exit Ticket")
    add_block(slide, 0.9, 1.25, 11.5, 4.1, "Answer before you leave", ["What did you finish?", "What did AI help with?", "What did you decide yourself?", "What would you improve next time?"], body_size=27)
    add_text(slide, 0.9, 6.0, 11.5, 0.5, "Save your work before you close your Chromebook.", size=26, bold=True, color=ACCENT4, align=PP_ALIGN.CENTER)

    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
