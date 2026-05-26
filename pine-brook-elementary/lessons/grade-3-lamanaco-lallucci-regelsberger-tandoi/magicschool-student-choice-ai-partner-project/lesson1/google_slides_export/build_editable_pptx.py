from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


@dataclass(frozen=True)
class Theme:
    accent1: RGBColor = RGBColor.from_string("0F766E")
    accent2: RGBColor = RGBColor.from_string("99F6E4")
    accent3: RGBColor = RGBColor.from_string("F59E0B")
    accent4: RGBColor = RGBColor.from_string("134E4A")
    paper: RGBColor = RGBColor.from_string("F7FEFC")
    ink: RGBColor = RGBColor.from_string("1F2937")


THEME = Theme()

# IMPORTANT: Google Slides cannot run “copy-to-clipboard” code inside a slide.
# The reliable workaround is: click a Copy chip -> open a web page -> click to copy there.
# Host `docs/lesson1-prompt-cards-copy.html` somewhere HTTPS-accessible (e.g., GitHub Pages)
# and set this URL to the hosted location.
COPY_PAGE_URL_BASE = "https://pzg8794.github.io/teaching-placement/lesson1-prompt-cards-copy.html"


def _add_rect(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    fill: RGBColor | None = None,
    line: RGBColor | None = None,
    line_width_pt: float | None = None,
    rounded: bool = False,
):
    shape_type = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if rounded else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shp = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill

    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        if line_width_pt is not None:
            shp.line.width = Pt(line_width_pt)
    return shp


def _add_textbox(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    text: str,
    *,
    size_pt: float,
    bold: bool = False,
    color: RGBColor = THEME.ink,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    font: str = "Arial",
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def _add_title(slide, title: str):
    # Simple consistent title style
    _add_textbox(slide, 0.7, 0.35, 12.0, 0.6, title, size_pt=34, bold=True, color=THEME.ink)


def _add_block(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    header: str | None = None,
    *,
    fill: RGBColor = RGBColor(255, 255, 255),
    border: RGBColor = THEME.accent4,
):
    rect = _add_rect(slide, x, y, w, h, fill=fill, line=border, line_width_pt=1.0, rounded=True)
    if header:
        _add_textbox(slide, x + 0.25, y + 0.15, w - 0.5, 0.35, header, size_pt=18, bold=True, color=THEME.accent4)
    return rect


def _add_bullet_list(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    items: list[str],
    *,
    size_pt: float = 20,
    bullet_char: str = "▶",
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True

    # First paragraph already exists
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT

        r1 = p.add_run()
        r1.text = f"{bullet_char} "
        r1.font.name = "Arial"
        r1.font.size = Pt(size_pt)
        r1.font.bold = True
        r1.font.color.rgb = THEME.accent1

        r2 = p.add_run()
        r2.text = item
        r2.font.name = "Arial"
        r2.font.size = Pt(size_pt)
        r2.font.bold = False
        r2.font.color.rgb = THEME.ink

    return box


def _add_label_box(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    text: str,
    *,
    fill: RGBColor,
    border: RGBColor,
):
    rect = _add_rect(slide, x, y, w, h, fill=fill, line=border, line_width_pt=1.0, rounded=True)
    _add_textbox(slide, x, y + (h - 0.4) / 2, w, 0.4, text, size_pt=18, bold=True, color=THEME.ink, align=PP_ALIGN.CENTER)
    return rect


def _add_copy_chip(slide, x: float, y: float, prompt_id: str, *, w: float = 0.85, h: float = 0.34):
    chip = _add_rect(
        slide,
        x,
        y,
        w,
        h,
        fill=RGBColor.from_string("CCFBF1"),
        line=THEME.accent1,
        line_width_pt=0.8,
        rounded=True,
    )
    tf = chip.text_frame
    tf.clear()
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "Copy"
    r.font.name = "Arial"
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = THEME.accent4

    chip.click_action.hyperlink.address = f"{COPY_PAGE_URL_BASE}#{prompt_id}"
    return chip


def build_deck(out_path: Path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # --- Slide 1: Title ---
    s = prs.slides.add_slide(blank)
    _add_rect(s, 0, 0, 13.333, 7.5, fill=THEME.paper, line=None)
    _add_rect(s, 0.55, 0.55, 12.25, 6.4, fill=RGBColor.from_string("E6FFFB"), line=None, rounded=True)
    _add_rect(s, 0.55, 0.65, 12.25, 0.08, fill=THEME.accent2, line=None)
    _add_rect(s, 0.55, 0.75, 12.25, 0.08, fill=THEME.accent3, line=None)

    _add_textbox(s, 0.85, 1.2, 7.2, 0.9, "Lesson 1: Brainstorm and Plan", size_pt=44, bold=True, color=THEME.ink)
    _add_textbox(s, 0.85, 2.0, 7.2, 0.6, "Pick a topic. Check it. Plan it.", size_pt=28, bold=False, color=THEME.accent4)

    ai_box = _add_rect(s, 0.85, 2.75, 6.8, 0.85, fill=RGBColor(255, 255, 255), line=THEME.accent1, line_width_pt=1.2, rounded=True)
    _ = ai_box
    _add_textbox(s, 0.95, 2.95, 6.6, 0.55, "AI helps us think. We stay in charge.", size_pt=26, bold=True, color=THEME.ink, align=PP_ALIGN.CENTER)

    step_w = 2.05
    step_h = 0.85
    step_y = 3.85
    _add_label_box(s, 0.85, step_y, step_w, step_h, "1\nIdea", fill=RGBColor.from_string("FDE68A"), border=THEME.accent3)
    _add_label_box(s, 0.85 + step_w + 0.25, step_y, step_w, step_h, "2\nCheck", fill=RGBColor.from_string("CCFBF1"), border=THEME.accent1)
    _add_label_box(s, 0.85 + 2 * (step_w + 0.25), step_y, step_w, step_h, "3\nGuide", fill=RGBColor.from_string("D1FAE5"), border=THEME.accent1)

    # Right-side simple mock device
    _add_rect(s, 9.1, 1.45, 3.7, 4.3, fill=RGBColor.from_string("CCFBF1"), line=THEME.accent4, line_width_pt=1.0, rounded=True)
    _add_rect(s, 9.5, 1.85, 2.9, 3.5, fill=RGBColor(255, 255, 255), line=THEME.accent1, line_width_pt=1.0, rounded=True)
    _add_rect(s, 9.75, 4.7, 2.4, 0.55, fill=RGBColor.from_string("99F6E4"), line=THEME.accent4, line_width_pt=1.0, rounded=True)
    _add_textbox(s, 9.75, 4.82, 2.4, 0.35, "topic + plan", size_pt=14, bold=True, color=THEME.ink, align=PP_ALIGN.CENTER)

    # --- Slide 2: Today's Mission ---
    s = prs.slides.add_slide(blank)
    _add_title(s, "Today's Mission")
    _add_block(s, 0.8, 1.25, 11.7, 3.1, "By the end of Lesson 1, I will have:", fill=RGBColor(255, 255, 255), border=THEME.accent4)
    _add_bullet_list(
        s,
        1.1,
        1.85,
        11.0,
        2.35,
        [
            "one topic I care about",
            "one reason it connects to computer science",
            "one project type",
            "a guide with 3 to 5 parts, sections, or questions",
        ],
        size_pt=22,
    )

    # Topic / CS Link / Guide chips
    chip_y = 4.65
    chip_w = 3.2
    chip_h = 0.7
    _add_label_box(s, 2.2, chip_y, chip_w, chip_h, "Topic", fill=RGBColor.from_string("FDE68A"), border=THEME.accent3)
    _add_label_box(s, 5.05, chip_y, chip_w, chip_h, "CS Link", fill=RGBColor.from_string("CCFBF1"), border=THEME.accent1)
    _add_label_box(s, 7.9, chip_y, chip_w, chip_h, "Guide", fill=RGBColor.from_string("D1FAE5"), border=THEME.accent1)

    # --- Slide 3: Our AI Rules ---
    s = prs.slides.add_slide(blank)
    _add_title(s, "Our AI Rules")
    _add_bullet_list(
        s,
        0.9,
        1.35,
        7.0,
        5.6,
        [
            "Do not type names or private information.",
            "AI is a helper, not the boss.",
            "AI can be wrong.",
            "We choose what to keep or change.",
            "The teacher can see class AI work.",
        ],
        size_pt=22,
    )
    _add_block(s, 8.2, 1.35, 4.2, 5.6, None, fill=RGBColor(255, 255, 255), border=THEME.accent4)
    _add_label_box(s, 8.55, 1.75, 3.5, 0.8, "Keep school rules in the chat.", fill=RGBColor.from_string("D1FAE5"), border=THEME.accent1)
    _add_label_box(s, 8.55, 3.15, 3.5, 0.9, "Read the answer.\nThink. Then decide.", fill=RGBColor.from_string("CCFBF1"), border=THEME.accent1)
    _add_label_box(s, 8.55, 4.75, 3.5, 0.8, "Never copy without checking.", fill=RGBColor.from_string("FDE68A"), border=THEME.accent3)

    # --- Slide 4: Use the Right Tool ---
    s = prs.slides.add_slide(blank)
    _add_title(s, "Use the Right Tool")
    _add_block(s, 0.8, 1.5, 5.9, 3.0, "Idea Generator", fill=RGBColor(255, 255, 255), border=THEME.accent4)
    _add_bullet_list(s, 1.15, 2.1, 5.2, 2.2, ["Start here", "Get possible topic ideas", "Connect your interests to CS"], size_pt=22)
    _add_block(s, 6.65, 1.5, 5.9, 3.0, "AI Assistant", fill=RGBColor(255, 255, 255), border=THEME.accent4)
    _add_bullet_list(
        s,
        7.0,
        2.1,
        5.2,
        2.2,
        [
            "Use after you have an idea",
            "Check the CS connection",
            "Make the topic smaller",
            "Build a guide with 3 to 5 parts",
        ],
        size_pt=22,
    )
    _add_textbox(
        s,
        0.8,
        4.85,
        12.0,
        0.6,
        "Idea Generator first. AI Assistant second.",
        size_pt=28,
        bold=True,
        color=THEME.ink,
        align=PP_ALIGN.CENTER,
    )

    # --- Slide 5: Our Workflow ---
    s = prs.slides.add_slide(blank)
    _add_title(s, "Our Workflow")
    step_y = 1.75
    w = 2.75
    h = 1.55
    xs = [0.9, 3.9, 6.9, 9.9]
    fills = [RGBColor.from_string("FDE68A"), RGBColor.from_string("CCFBF1"), RGBColor.from_string("D1FAE5"), RGBColor.from_string("FDE68A")]
    labels = [
        "1\nPick an interest",
        "2\nUse Idea Generator",
        "3\nUse AI Assistant",
        "4\nRecord your guide",
    ]
    for x, fill, lab in zip(xs, fills, labels, strict=True):
        _add_label_box(s, x, step_y, w, h, lab, fill=fill, border=THEME.accent4)

    _add_block(s, 0.9, 3.75, 12.0, 1.1, "Remember", fill=RGBColor(255, 255, 255), border=THEME.accent4)
    _add_textbox(
        s,
        1.15,
        4.1,
        11.5,
        0.6,
        "You do not need a finished project today. You need a strong topic and a strong guide.",
        size_pt=22,
        bold=False,
        color=THEME.ink,
    )

    # --- Slide 6: Prompt Cards - Idea Generator ---
    s = prs.slides.add_slide(blank)
    _add_title(s, "Prompt Cards: Idea Generator")
    _add_block(s, 0.8, 1.45, 12.0, 3.2, "Use Idea Generator for topic ideas", fill=RGBColor(255, 255, 255), border=THEME.accent4)
    _add_bullet_list(
        s,
        1.05,
        2.1,
        11.5,
        2.4,
        [
            "My interest is ____. Give me 4 computer-science project ideas for grades 4-5.",
            "Help me turn my interest in ____ into a computer-science project topic.",
        ],
        size_pt=20,
    )
    _add_copy_chip(s, 11.75, 2.1, "ig1")
    _add_copy_chip(s, 11.75, 2.72, "ig2")
    _add_textbox(s, 0.8, 4.95, 12.0, 0.6, "Pick 1 or 2 ideas worth checking.", size_pt=28, bold=True, align=PP_ALIGN.CENTER)

    # --- Slide 7: Prompt Cards - AI Assistant ---
    s = prs.slides.add_slide(blank)
    _add_title(s, "Prompt Cards: AI Assistant")
    _add_block(s, 0.8, 1.45, 12.0, 3.6, "Use AI Assistant for checking and planning", fill=RGBColor(255, 255, 255), border=THEME.accent4)
    _add_bullet_list(
        s,
        1.05,
        2.1,
        11.5,
        2.8,
        [
            "My topic is ____. Is this really connected to computer science? Explain how.",
            "My topic is ____. Help me make it smaller and clearer for a short project.",
            "My topic is ____. Give me a high-level guide with 3 to 5 parts, sections, or questions I could use for my project.",
        ],
        size_pt=18,
    )
    _add_copy_chip(s, 11.75, 2.08, "aa1")
    _add_copy_chip(s, 11.75, 2.66, "aa2")
    _add_copy_chip(s, 11.75, 3.24, "aa3")
    _add_textbox(s, 0.8, 5.25, 12.0, 0.6, "Choose a guide you agree with.", size_pt=28, bold=True, align=PP_ALIGN.CENTER)

    # --- Slide 8: What Not To Do ---
    s = prs.slides.add_slide(blank)
    _add_title(s, "What Not To Do")
    _add_block(s, 0.8, 1.55, 6.1, 4.8, "Do", fill=RGBColor(255, 255, 255), border=THEME.accent4)
    _add_bullet_list(
        s,
        1.05,
        2.25,
        5.6,
        3.9,
        [
            "read the answer",
            "check if it fits your goal",
            "pick the parts you agree with",
            "ask for a better version if needed",
        ],
        size_pt=22,
    )
    _add_block(s, 6.9, 1.55, 6.1, 4.8, "Do Not", fill=RGBColor(255, 255, 255), border=THEME.accent4)
    _add_bullet_list(
        s,
        7.15,
        2.25,
        5.6,
        3.9,
        [
            "copy the whole answer",
            "accept the first answer right away",
            "keep a topic that is too long/complex",
            "choose a topic with no CS connection",
        ],
        size_pt=22,
    )

    # --- Slide 9: Planning Sheet ---
    s = prs.slides.add_slide(blank)
    _add_title(s, "What Goes on Your Planning Sheet?")
    _add_bullet_list(
        s,
        0.95,
        1.55,
        12.2,
        2.3,
        [
            "My topic:",
            "My CS connection:",
            "My project type:",
            "My guide with 3 to 5 parts, sections, or questions:",
        ],
        size_pt=24,
    )
    _add_block(s, 0.95, 3.95, 12.2, 1.05, "Optional AI helper", fill=RGBColor(255, 255, 255), border=THEME.accent3)
    _add_copy_chip(s, 11.85, 4.12, "ps1")
    _add_textbox(
        s,
        1.2,
        4.35,
        11.7,
        0.55,
        "My topic is ____. My interest is ____. Help me fill out my planning sheet in the right format.",
        size_pt=18,
        bold=False,
        color=THEME.ink,
    )
    _add_block(s, 0.95, 5.2, 12.2, 1.25, "Teacher check", fill=RGBColor(255, 255, 255), border=THEME.accent4)
    _add_textbox(s, 1.2, 5.65, 11.7, 0.6, "Before Lesson 2, your topic and guide must be approved.", size_pt=24, bold=True)

    # --- Slide 10: If You Get Stuck ---
    s = prs.slides.add_slide(blank)
    _add_title(s, "If You Get Stuck")
    _add_block(s, 0.8, 1.45, 7.0, 5.4, "Troubleshoot in order", fill=RGBColor(255, 255, 255), border=THEME.accent4)
    _add_bullet_list(
        s,
        1.05,
        2.1,
        6.5,
        4.6,
        [
            "Name the problem.",
            "Try one small change.",
            "Use the matching prompt.",
            "Read the answer. Keep only what helps you.",
            "Ask your partner or teacher.",
        ],
        size_pt=22,
    )
    _add_block(s, 8.1, 1.45, 5.1, 3.9, "Problem key", fill=RGBColor(255, 255, 255), border=THEME.accent4)
    key_items = [
        ("No topic yet", RGBColor.from_string("FDE68A"), THEME.accent3),
        ("Too big", RGBColor.from_string("D1FAE5"), THEME.accent1),
        ("Too simple", RGBColor.from_string("CCFBF1"), THEME.accent1),
        ("Don't like ideas", RGBColor.from_string("E6FFFB"), THEME.accent4),
        ("Not sure it is CS", RGBColor.from_string("CCFBF1"), THEME.accent4),
    ]
    ky = 2.05
    for label, fill, border in key_items:
        _add_label_box(s, 8.45, ky, 4.4, 0.55, label, fill=fill, border=border)
        ky += 0.7

    # --- Slide 11: Copy-Paste Prompts ---
    s = prs.slides.add_slide(blank)
    _add_title(s, "If You Get Stuck: Copy-Paste Prompts")
    _add_block(s, 0.8, 1.45, 12.0, 5.9, "Copy one prompt and try it", fill=RGBColor(255, 255, 255), border=THEME.accent4)

    prompts = [
        ("ts1", "No topic yet", "I do not know what topic to choose yet. My interests are: ____. Give me 8 different computer-science project topic ideas."),
        ("ts2", "Still stuck", "I do not know what I like yet. Ask me 5 quick questions. Then give me 6 topic ideas connected to computer science."),
        ("ts3", "Too big", "My topic is ____. It feels too big or too complex. Help me make it smaller. Give me 3 smaller topic options."),
        ("ts4", "Too simple", "My topic is ____. It feels too simple. Help me make it more interesting (but still doable). Give me 3 upgraded topic versions."),
        ("ts5", "Don't like ideas", "I do not like these ideas: ____. Ask me 3 questions. Then give me 8 new topic ideas. Do NOT repeat the old ideas."),
    ]

    # Two-column layout for readability
    left = prompts[:3]
    right = prompts[3:]

    def add_prompt_column(col_x: float, col_prompts: list[tuple[str, str, str]]):
        y0 = 2.05
        for pid, title, text in col_prompts:
            _add_rect(s, col_x, y0, 5.75, 1.55, fill=RGBColor.from_string("F9FAFB"), line=THEME.accent4, line_width_pt=0.8, rounded=True)
            _add_copy_chip(s, col_x + 5.75 - 0.95, y0 + 0.14, pid)
            _add_textbox(s, col_x + 0.2, y0 + 0.15, 5.35, 0.3, title, size_pt=16, bold=True, color=THEME.accent4)
            _add_textbox(s, col_x + 0.2, y0 + 0.45, 5.35, 1.0, text, size_pt=14, bold=False, color=THEME.ink)
            y0 += 1.75

    add_prompt_column(1.05, left)
    add_prompt_column(7.05, right)

    # --- Slide 12: Ready for Lesson 2? ---
    s = prs.slides.add_slide(blank)
    _add_title(s, "Ready for Lesson 2?")
    y = 1.55
    for text, fill, border in [
        ("I have one topic.", RGBColor.from_string("D1FAE5"), THEME.accent1),
        ("I can explain the CS connection.", RGBColor.from_string("CCFBF1"), THEME.accent1),
        ("I have a guide with 3 to 5 parts.", RGBColor.from_string("FDE68A"), THEME.accent3),
    ]:
        _add_label_box(s, 1.0, y, 11.3, 0.8, text, fill=fill, border=border)
        y += 0.95

    _add_block(s, 1.0, 4.65, 11.3, 1.55, "Quick share", fill=RGBColor(255, 255, 255), border=THEME.accent4)
    _add_textbox(
        s,
        1.25,
        5.15,
        10.9,
        0.85,
        "My topic is…  My project will be…  My guide will help me by…",
        size_pt=26,
        bold=False,
        color=THEME.ink,
    )

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out_path)


if __name__ == "__main__":
    here = Path(__file__).resolve().parent
    out = here / "lesson1-brainstorm-and-plan-slides-editable-with-copy-links.pptx"
    build_deck(out)
    print(f"Wrote: {out}")
